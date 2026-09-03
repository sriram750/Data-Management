import io
import json
import openpyxl
import pytest
from httpx import AsyncClient


def create_sample_excel_bytes() -> bytes:
    """Generates an in-memory test .xlsx file with actual data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees_Q3"
    
    headers = ["Emp_ID", "Full_Name", "Department", "Email", "Salary"]
    ws.append(headers)
    
    rows = [
        ["E101", "John Doe", "Engineering", "john.doe@company.org", 85000],
        ["E102", "Jane Smith", "Finance", "jane.smith@company.org", 92000],
        ["E103", "Bob Wilson", "HR", "bob.wilson@company.org", 65000],
    ]
    for r in rows:
        ws.append(r)
        
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.asyncio
async def test_excel_preview_validate_and_import(client: AsyncClient, super_admin_auth):
    _, headers = super_admin_auth
    excel_bytes = create_sample_excel_bytes()

    # 1. Upload for Preview
    files = {"file": ("test_employees.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview_res = await client.post("/api/v1/imports/upload", files=files, headers=headers)
    assert preview_res.status_code == 200
    preview = preview_res.json()
    assert preview["total_rows"] == 3
    assert preview["total_columns"] == 5
    assert len(preview["detected_columns"]) == 5
    file_token = preview["file_token"]
    sheet_name = preview["selected_sheet"]

    # 2. Validate Rows
    cols_config = [
        {"source_column": "Emp_ID", "target_column": "emp_id", "data_type": "TEXT", "is_required": True},
        {"source_column": "Full_Name", "target_column": "full_name", "data_type": "TEXT", "is_required": True},
        {"source_column": "Department", "target_column": "department", "data_type": "DROPDOWN"},
        {"source_column": "Email", "target_column": "email", "data_type": "EMAIL"},
        {"source_column": "Salary", "target_column": "salary", "data_type": "CURRENCY"},
    ]
    val_res = await client.post(
        "/api/v1/imports/validate",
        data={"file_token": file_token, "sheet_name": sheet_name, "columns_json": json.dumps(cols_config)},
        headers=headers,
    )
    assert val_res.status_code == 200
    val_data = val_res.json()
    assert val_data["total_rows"] == 3
    assert val_data["valid_rows"] == 3
    assert val_data["error_rows"] == 0
    assert val_data["is_valid"] is True

    # 3. Execute Import (Create New Table)
    exec_payload = {
        "file_token": file_token,
        "sheet_name": sheet_name,
        "mode": "INSERT_NEW_TABLE",
        "new_table_name": "q3_employees",
        "new_table_display_name": "Q3 Employees Imported",
        "new_table_description": "Imported from Q3 spreadsheet",
        "columns": cols_config,
    }
    exec_res = await client.post("/api/v1/imports/execute", json=exec_payload, headers=headers)
    assert exec_res.status_code == 200
    import_history = exec_res.json()
    assert import_history["imported_rows"] == 3
    assert import_history["error_rows"] == 0
    table_id = import_history["table_id"]

    # 4. Verify the newly created table contains the 3 imported records
    records_res = await client.get(f"/api/v1/tables/{table_id}/records", headers=headers)
    assert records_res.status_code == 200
    records_data = records_res.json()
    assert records_data["total"] == 3
    assert records_data["items"][0]["data"]["emp_id"] == "E101"
    assert records_data["items"][1]["data"]["full_name"] == "Jane Smith"

    # 5. Verify Import History endpoint lists this import
    hist_res = await client.get("/api/v1/imports/history", headers=headers)
    assert hist_res.status_code == 200
    assert len(hist_res.json()) >= 1


@pytest.mark.asyncio
async def test_excel_import_mixed_data_type(client: AsyncClient, super_admin_auth):
    _, headers = super_admin_auth
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MixedData"
    ws.append(["EmpId", "EmpName", "Extension", "2025-DEC"])
    ws.append(["101", "Alice", "123", 10.5])
    ws.append(["102", "Bob", "N@R", "#N/A"])
    ws.append(["103", "Charlie", "DFJR", "WFH/ Own Headphone"])
    ws.append(["104", "Dave", "456", "69012405500010-"])
    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    # 1. Upload for Preview
    files = {"file": ("mixed_test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview_res = await client.post("/api/v1/imports/upload", files=files, headers=headers)
    assert preview_res.status_code == 200
    preview = preview_res.json()
    file_token = preview["file_token"]
    sheet_name = preview["selected_sheet"]

    # Inferred types should detect MIXED for Extension and 2025-DEC
    col_map = {c["original_name"]: c["suggested_type"] for c in preview["detected_columns"]}
    assert col_map["Extension"] == "MIXED"
    assert col_map["2025-DEC"] == "MIXED"

    # 2. Validate with MIXED
    cols_config = [
        {"source_column": "EmpId", "target_column": "empid", "data_type": "TEXT"},
        {"source_column": "EmpName", "target_column": "empname", "data_type": "TEXT"},
        {"source_column": "Extension", "target_column": "extension", "data_type": "MIXED"},
        {"source_column": "2025-DEC", "target_column": "2025_dec", "data_type": "MIXED"},
    ]
    val_res = await client.post(
        "/api/v1/imports/validate",
        data={"file_token": file_token, "sheet_name": sheet_name, "columns_json": json.dumps(cols_config)},
        headers=headers,
    )
    assert val_res.status_code == 200
    val_data = val_res.json()
    assert val_data["error_rows"] == 0
    assert val_data["valid_rows"] == 4
    assert val_data["is_valid"] is True

    # 3. Execute Import
    exec_payload = {
        "file_token": file_token,
        "sheet_name": sheet_name,
        "mode": "INSERT_NEW_TABLE",
        "new_table_name": "mixed_table_test",
        "new_table_display_name": "Mixed Table Test",
        "columns": cols_config,
    }
    exec_res = await client.post("/api/v1/imports/execute", json=exec_payload, headers=headers)
    assert exec_res.status_code == 200
    table_id = exec_res.json()["table_id"]

    # Verify records
    records_res = await client.get(f"/api/v1/tables/{table_id}/records", headers=headers)
    assert records_res.status_code == 200
    items = records_res.json()["items"]
    assert len(items) == 4
    assert items[1]["data"]["2025_dec"] == "#N/A"
    assert items[2]["data"]["2025_dec"] == "WFH/ Own Headphone"
    assert items[3]["data"]["2025_dec"] == "69012405500010-"


@pytest.mark.asyncio
async def test_excel_multi_sheet_batch_import(client: AsyncClient, super_admin_auth):
    _, headers = super_admin_auth
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Employees"
    ws1.append(["Emp_ID", "Name", "Role"])
    ws1.append(["E001", "Alice", "Engineer"])
    ws1.append(["E002", "Charlie", "Designer"])

    ws2 = wb.create_sheet(title="Exit_Staff")
    ws2.append(["Exit_ID", "Emp_ID", "Reason"])
    ws2.append(["X001", "E001", "Relocation"])

    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    # Upload for preview
    files = {"file": ("multi_sheet_test.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview_res = await client.post("/api/v1/imports/upload", files=files, headers=headers)
    assert preview_res.status_code == 200
    preview = preview_res.json()
    assert len(preview["sheets"]) == 2
    assert "Employees" in preview["sheets"]
    assert "Exit_Staff" in preview["sheets"]

    # Execute Batch Import for both sheets
    batch_payload = {
        "file_token": preview["file_token"],
        "sheets": [
            {
                "sheet_name": "Employees",
                "table_display_name": "Active Employees",
                "table_name": "active_employees_batch",
            },
            {
                "sheet_name": "Exit_Staff",
                "table_display_name": "Exit Staff List",
                "table_name": "exit_staff_batch",
            },
        ],
    }
    batch_res = await client.post("/api/v1/imports/batch-execute", json=batch_payload, headers=headers)
    assert batch_res.status_code == 200
    data = batch_res.json()
    assert data["total_sheets_processed"] == 2
    assert len(data["successful_tables"]) == 2
    assert len(data["failed_sheets"]) == 0

    # Verify both tables exist and have records
    emp_table = next(t for t in data["successful_tables"] if t["sheet_name"] == "Employees")
    exit_table = next(t for t in data["successful_tables"] if t["sheet_name"] == "Exit_Staff")
    assert emp_table["imported_rows"] == 2
    assert exit_table["imported_rows"] == 1

